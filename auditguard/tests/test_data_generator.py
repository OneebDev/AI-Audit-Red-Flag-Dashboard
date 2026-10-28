"""Tests for the reproducible Karachi demo dataset and its workbooks."""

from __future__ import annotations

from collections import Counter

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

import data_generator
from data_generator import (
    ACCOUNTS,
    APPROVAL_LIMIT,
    BENFORD_VENDOR,
    BENFORD_VENDOR_ROWS,
    BLIND_DATASET_NAME,
    CLEAN_ROWS,
    DATASET_NAME,
    DEFAULT_ENTRIES,
    DUPLICATE_PAIRS,
    EMPLOYEES,
    MULTIRULE_COUNT,
    RENT_AMOUNT,
    RENT_VENDOR,
    RENT_VENDOR_ROWS,
    ROUND_AMOUNTS,
    VENDOR_ACCOUNTS,
    about_frame,
    generate_ledger,
    holidays_frame,
    write_datasets,
)
from detection import (
    FLAG_COLUMNS,
    GROUND_TRUTH_COLUMN,
    REQUIRED_COLUMNS,
)
from detection import ALL_LABELS, CLEAN_LABEL, MULTIRULE_LABEL


def _leading_digits(values: pd.Series) -> pd.Series:
    usable = pd.to_numeric(values, errors="coerce").abs()
    usable = usable[usable >= 1]
    exponent = np.floor(np.log10(usable.to_numpy(dtype=float)))
    return pd.Series(
        np.floor(usable.to_numpy(dtype=float) / np.power(10.0, exponent)).astype(int),
        index=usable.index,
    )


@pytest.fixture(scope="module")
def ledger() -> pd.DataFrame:
    return generate_ledger()


def test_ledger_has_the_fixed_schema_and_size(ledger: pd.DataFrame) -> None:
    assert list(ledger.columns) == list(REQUIRED_COLUMNS) + [GROUND_TRUTH_COLUMN]
    assert len(ledger) == DEFAULT_ENTRIES == 606
    assert ledger["EntryID"].is_unique
    assert ledger["EntryID"].iloc[0] == "ENT-0001"
    assert ledger["EntryID"].iloc[-1] == "ENT-0606"


def test_ground_truth_counts(ledger: pd.DataFrame) -> None:
    counts = Counter(ledger[GROUND_TRUTH_COLUMN])

    assert counts[CLEAN_LABEL] == CLEAN_ROWS == 549
    assert counts["Duplicate"] == 12
    assert counts["RoundAmount"] == 5
    assert counts["Weekend"] == 5
    assert counts["JustBelowLimit"] == 6
    assert counts["SelfApproval"] == 5
    assert counts["BenfordManip"] == BENFORD_VENDOR_ROWS == 18
    assert counts[MULTIRULE_LABEL] == MULTIRULE_COUNT == 6
    assert sum(counts[label] for label in ALL_LABELS if label != CLEAN_LABEL) == 57


def test_reference_data_sizes(ledger: pd.DataFrame) -> None:
    assert len(VENDOR_ACCOUNTS) == 25
    assert set(ledger["Vendor"]) == set(VENDOR_ACCOUNTS)
    assert set(ledger["Account"]) == set(ACCOUNTS)
    assert len(ACCOUNTS) == 10
    assert set(ledger["PostedBy"]) == set(EMPLOYEES)
    assert len(EMPLOYEES) == 6


def test_dates_cover_2025(ledger: pd.DataFrame) -> None:
    dates = pd.to_datetime(ledger["Date"])

    assert dates.dt.year.unique().tolist() == [2025]
    assert dates.min() >= pd.Timestamp("2025-01-01")
    assert dates.max() <= pd.Timestamp("2025-12-31")


def test_generation_is_deterministic() -> None:
    first = generate_ledger(seed=42)
    second = generate_ledger(seed=42)
    other = generate_ledger(seed=7)

    pd.testing.assert_frame_equal(first, second)
    assert not first["Amount"].equals(other["Amount"])


def test_clean_rows_respect_every_rule(ledger: pd.DataFrame, checked_demo: pd.DataFrame) -> None:
    clean = ledger[ledger[GROUND_TRUTH_COLUMN] == CLEAN_LABEL]

    assert clean["Date"].dt.dayofweek.lt(5).all()
    assert not set(clean["Date"].dt.normalize()) & set(data_generator.HOLIDAY_DATES_2025)
    assert (clean["PostedBy"].str.casefold() != clean["ApprovedBy"].str.casefold()).all()

    amounts = pd.to_numeric(clean["Amount"], errors="coerce")
    assert amounts.between(1_000.0, 380_000.0).all()
    round_mask = (amounts >= 10_000) & (np.mod(amounts, 1_000.0).abs() < 0.005)
    assert not round_mask.any()

    checked_clean = checked_demo[checked_demo[GROUND_TRUTH_COLUMN] == CLEAN_LABEL]
    assert checked_clean["flag_count"].sum() == 0
    assert checked_clean[list(FLAG_COLUMNS)].to_numpy().sum() == 0


def test_clean_rows_avoid_repeats_within_seven_days(ledger: pd.DataFrame) -> None:
    clean = ledger[ledger[GROUND_TRUTH_COLUMN] == CLEAN_LABEL]
    work = clean[["Vendor", "Amount", "Date"]].copy()
    work["Amount"] = work["Amount"].round(2)
    work = work.sort_values(["Vendor", "Amount", "Date"])

    for _, group in work.groupby(["Vendor", "Amount"], sort=False):
        if len(group) < 2:
            continue
        gaps = group["Date"].diff().dt.days.dropna()
        assert (gaps > 7).all()


def test_rent_vendor_is_a_fixed_monthly_series(ledger: pd.DataFrame) -> None:
    rent = ledger[ledger["Vendor"] == RENT_VENDOR]

    assert len(rent) == RENT_VENDOR_ROWS == 12
    assert rent["Amount"].eq(RENT_AMOUNT).all()
    assert rent["Date"].dt.month.nunique() == 12
    assert rent["Date"].dt.dayofweek.lt(5).all()
    assert rent["Account"].eq("Rent").all()


def test_benford_vendor_uses_fabricated_leading_digits(ledger: pd.DataFrame) -> None:
    benford = ledger[ledger[GROUND_TRUTH_COLUMN] == "BenfordManip"]

    assert len(benford) == 18
    assert benford["Vendor"].eq(BENFORD_VENDOR).all()
    assert set(_leading_digits(benford["Amount"])) <= {5, 6, 7, 8, 9}


def test_duplicates_are_six_labelled_pairs(ledger: pd.DataFrame) -> None:
    duplicates = ledger[ledger[GROUND_TRUTH_COLUMN] == "Duplicate"]

    assert len(duplicates) == DUPLICATE_PAIRS * 2 == 12

    paired = duplicates.groupby(["Vendor", "Amount"])
    assert len(paired) == 6
    for _, group in paired:
        assert len(group) == 2
        gap = (group["Date"].max() - group["Date"].min()).days
        assert 1 <= gap <= 7
        assert group["Date"].dt.dayofweek.lt(5).all()

    descriptions = duplicates["Description"]
    assert int(descriptions.str.contains("INV-2025-0142", regex=False).sum()) == 2
    assert descriptions.nunique() == 6


def test_single_rule_anomalies_have_the_expected_shape(ledger: pd.DataFrame) -> None:
    round_rows = ledger[ledger[GROUND_TRUTH_COLUMN] == "RoundAmount"]
    assert len(round_rows) == 5
    assert set(round_rows["Amount"]) == set(ROUND_AMOUNTS)
    assert round_rows["Date"].dt.dayofweek.lt(5).all()

    weekend = ledger[ledger[GROUND_TRUTH_COLUMN] == "Weekend"]
    assert len(weekend) == 5
    assert (weekend["Date"].dt.dayofweek >= 5).all()
    assert not set(weekend["Date"].dt.normalize()) & set(data_generator.HOLIDAY_DATES_2025)

    threshold = ledger[ledger[GROUND_TRUTH_COLUMN] == "JustBelowLimit"]
    assert len(threshold) == 6
    assert threshold["Amount"].between(APPROVAL_LIMIT * 0.95, APPROVAL_LIMIT * 0.999).all()
    assert threshold["Date"].dt.dayofweek.lt(5).all()

    self_approval = ledger[ledger[GROUND_TRUTH_COLUMN] == "SelfApproval"]
    assert len(self_approval) == 5
    assert (
        self_approval["PostedBy"].str.casefold()
        == self_approval["ApprovedBy"].str.casefold()
    ).all()


def test_single_rule_rows_trigger_exactly_one_flag(checked_demo: pd.DataFrame) -> None:
    labels = [
        "Duplicate",
        "RoundAmount",
        "Weekend",
        "JustBelowLimit",
        "SelfApproval",
        "BenfordManip",
    ]

    for label in labels:
        rows = checked_demo[checked_demo[GROUND_TRUTH_COLUMN] == label]
        assert rows["flag_count"].eq(1).all(), label


def test_multi_rule_rows_trip_two_or_three_rules(ledger: pd.DataFrame, checked_demo: pd.DataFrame) -> None:
    multi = ledger[ledger[GROUND_TRUTH_COLUMN] == MULTIRULE_LABEL]

    assert len(multi) == 6
    assert multi["PostedBy"].eq(multi["ApprovedBy"]).all()
    assert multi["EntryID"].is_unique

    checked = checked_demo[checked_demo[GROUND_TRUTH_COLUMN] == MULTIRULE_LABEL]
    assert checked["flag_count"].between(2, 3).all()
    assert (checked["flag_count"] == 3).sum() == 4
    assert checked["risk_level"].value_counts().to_dict() == {"High": 4, "Critical": 2}


def test_small_ledgers_skip_multi_rule_rows() -> None:
    frame = generate_ledger(n=150)

    assert len(frame) == 150
    assert MULTIRULE_LABEL not in set(frame[GROUND_TRUTH_COLUMN])
    assert frame[GROUND_TRUTH_COLUMN].value_counts()[CLEAN_LABEL] == 150 - 51


def test_parameter_validation() -> None:
    with pytest.raises(ValueError):
        generate_ledger(n=50)
    with pytest.raises(ValueError):
        generate_ledger(approval_limit=0)


def test_holidays_frame_lists_fixed_and_moon_sighting_dates() -> None:
    frame = holidays_frame()

    assert list(frame.columns) == ["Date", "Holiday", "Type"]
    assert len(frame) == 11
    fixed = frame[frame["Type"] == "Pakistan public holiday"]
    moon = frame[frame["Type"].str.contains("Moon-sighting dependent")]
    assert len(fixed) == 7
    assert len(moon) == 4
    assert pd.Timestamp("2025-02-05") in set(fixed["Date"])
    assert "Kashmir Day" in set(fixed["Holiday"])


def test_about_frame_reports_dataset_counts(ledger: pd.DataFrame) -> None:
    about = about_frame(ledger)
    values = dict(zip(about["Item"], about["Value"]))

    assert values["Total rows"] == 606
    assert values["Clean rows"] == 549
    assert values["Duplicate rows (6 pairs)"] == 12
    assert values["RoundAmount rows"] == 5
    assert values["Weekend rows"] == 5
    assert values["JustBelowLimit rows"] == 6
    assert values["SelfApproval rows"] == 5
    assert values["BenfordManip rows (Madina Traders)"] == 18
    assert values["MultiRule rows"] == 6
    assert values["Approval limit"] == "Rs. 500,000"
    assert values["Currency"] == "Pakistani Rupee (PKR, Rs.)"
    assert "ISA 240" in values["Note"]


def test_write_datasets_creates_labelled_and_blind_workbooks(tmp_path) -> None:
    labelled_path, blind_path = write_datasets(tmp_path)

    assert labelled_path.name == DATASET_NAME
    assert blind_path.name == BLIND_DATASET_NAME
    assert labelled_path.exists() and blind_path.exists()

    labelled = load_workbook(labelled_path, read_only=True)
    assert labelled.sheetnames == ["Ledger", "Pakistan_Holidays_2025", "About"]
    assert labelled["Ledger"].max_row == 607
    assert labelled["Ledger"].max_column == 9
    labelled.close()

    blind = load_workbook(blind_path, read_only=True)
    assert blind.sheetnames == ["Ledger"]
    assert blind["Ledger"].max_row == 607
    assert blind["Ledger"].max_column == 8
    header = [cell.value for cell in next(blind["Ledger"].iter_rows(max_row=1))]
    assert GROUND_TRUTH_COLUMN not in header
    blind.close()

    ledger = pd.read_excel(labelled_path, sheet_name="Ledger")
    assert len(ledger) == 606
    assert GROUND_TRUTH_COLUMN in ledger.columns

    about = pd.read_excel(labelled_path, sheet_name="About")
    assert "Total rows" in set(about["Item"])
