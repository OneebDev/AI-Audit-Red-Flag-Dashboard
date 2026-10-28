"""Tests for the Excel audit report, styling, and GroundTruth accuracy."""

from __future__ import annotations

from io import BytesIO

import pandas as pd
import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from detection import (
    FLAG_COLUMNS,
    GROUND_TRUTH_COLUMN,
    build_summary,
    run_all_checks,
)
from export import (
    GT_TO_FLAG,
    REPORT_COLUMNS,
    build_audit_report,
    compute_accuracy,
    dataframe_to_excel_bytes,
)

EXPECTED_SHEETS = [
    "Summary",
    "Flagged Entries",
    "Flag Breakdown",
    "AI Explanations",
    "Methodology",
]

RISK_FILLS = {
    "Low": "DCFCE7",
    "Medium": "FEF3C7",
    "High": "FFEDD5",
    "Critical": "FEE2E2",
}


@pytest.fixture(scope="module")
def report_artifacts(checked_demo, summary_demo):
    flagged = checked_demo[checked_demo["flag_count"] > 0].sort_values(
        "risk_score", ascending=False, kind="stable"
    )
    top_id = str(flagged.iloc[0]["EntryID"])
    explanations = {top_id: "Custom explanation for the top finding."}
    payload = build_audit_report(checked_demo, summary_demo, explanations)
    return payload, flagged, top_id


def test_report_has_all_five_sheets(report_artifacts) -> None:
    payload, _, _ = report_artifacts
    workbook = load_workbook(BytesIO(payload))

    assert payload[:2] == b"PK"
    assert workbook.sheetnames == EXPECTED_SHEETS
    assert [workbook[name].sheet_properties.tabColor.rgb for name in EXPECTED_SHEETS] == [
        "007C3AED",
        "0022D3EE",
        "0034D399",
        "00FBBF24",
        "00FB923C",
    ]


def test_summary_sheet_contains_metrics_formulas_and_charts(report_artifacts) -> None:
    payload, _, _ = report_artifacts
    worksheet = load_workbook(BytesIO(payload))["Summary"]

    assert worksheet["A1"].value == "AuditGuard Audit Report — Karachi"
    assert "Prepared in Karachi, Pakistan" in worksheet["A2"].value
    assert worksheet["B5"].value == 606
    assert worksheet["B6"].value == "=COUNTA('Flagged Entries'!$A$2:$A$100000)"
    assert worksheet["B7"].value == "=IFERROR(B6/B5,0)"
    assert str(worksheet["B8"].value).startswith("=SUMIF(")
    assert worksheet["A8"].value == "Amount at risk (Rs.)"
    assert "COUNTIF('Flagged Entries'!$J:$J" in str(worksheet["B15"].value)
    assert worksheet["B22"].value == "=COUNTIF('Flagged Entries'!$O:$O,TRUE)"
    assert worksheet["B28"].value == "=SUM(B22:B27)"
    assert len(worksheet._charts) == 2
    assert worksheet.sheet_view.showGridLines is False


def test_flag_formulas_follow_the_actual_column_layout(checked_demo, summary_demo) -> None:
    blind = checked_demo.drop(columns=[GROUND_TRUTH_COLUMN])
    payload = build_audit_report(blind, summary_demo)
    worksheet = load_workbook(BytesIO(payload))["Summary"]

    duplicate_letter = get_column_letter(
        [column for column in REPORT_COLUMNS if column != GROUND_TRUTH_COLUMN].index(
            "flag_duplicate"
        )
        + 1
    )
    assert duplicate_letter == "N"
    assert worksheet["B22"].value == (
        f"=COUNTIF('Flagged Entries'!${duplicate_letter}:${duplicate_letter},TRUE)"
    )


def test_flagged_sheet_is_filtered_formatted_and_counted(report_artifacts) -> None:
    payload, flagged, _ = report_artifacts
    worksheet = load_workbook(BytesIO(payload))["Flagged Entries"]

    assert worksheet.max_row == len(flagged) + 1 == 58
    assert worksheet.max_column == len(REPORT_COLUMNS) == 20
    assert worksheet["A1"].value == "Entry ID"
    assert worksheet["F1"].value == "Amount (Rs.)"
    assert worksheet["J1"].value == "Risk Level"
    assert worksheet.freeze_panes == "A2"
    assert worksheet.auto_filter.ref == "A1:T58"
    assert worksheet["B2"].number_format == "mm/dd/yyyy"
    assert worksheet["F2"].number_format == '"Rs. "#,##0.00'
    assert worksheet["I2"].number_format == "0"

    top_level = str(flagged.iloc[0]["risk_level"])
    assert worksheet["A2"].fill.start_color.rgb.endswith(RISK_FILLS[top_level])


def test_explanations_and_methodology_sheets(report_artifacts) -> None:
    payload, flagged, top_id = report_artifacts
    workbook = load_workbook(BytesIO(payload))

    explanations = workbook["AI Explanations"]
    assert explanations["A1"].value == "Entry ID"
    assert explanations["C1"].value == "Amount (Rs.)"
    assert explanations["E1"].value == "Explanation"
    assert explanations["A2"].value == top_id
    assert explanations["E2"].value == "Custom explanation for the top finding."

    methodology = workbook["Methodology"]
    rule_names = {
        methodology.cell(row=row, column=1).value for row in range(4, 10)
    }
    assert rule_names == {
        "Duplicate",
        "Round amount",
        "Weekend posting",
        "Just below limit",
        "Self-approval",
        "Benford deviation",
    }
    text = " ".join(
        str(cell.value)
        for row in methodology.iter_rows()
        for cell in row
        if cell.value is not None
    )
    assert "Rs. 500,000" in text
    assert "ISA 240" in text
    assert "Saturday + Sunday" in text


def test_zero_flag_ledger_still_builds_a_valid_report(make_ledger) -> None:
    checked = run_all_checks(
        make_ledger(
            [
                {"Vendor": "Vendor A", "Amount": 1_234.56},
                {"Vendor": "Vendor B", "Amount": 2_345.67},
            ]
        )
    )
    summary = build_summary(checked)
    payload = build_audit_report(checked, summary)
    workbook = load_workbook(BytesIO(payload))

    assert workbook.sheetnames == EXPECTED_SHEETS
    assert workbook["Flagged Entries"].max_row == 1
    assert len(workbook["Summary"]._charts) == 2
    assert len(payload) > 0


def test_formula_like_text_is_neutralized_in_reports(make_ledger) -> None:
    frame = make_ledger(
        [
            {
                "Vendor": "Vendor A",
                "Amount": 50_000.0,
                "Description": "=HYPERLINK(\"http://example.com\",\"click\")",
            },
            {"Vendor": "Vendor B", "Amount": 2_345.67},
        ]
    )
    checked = run_all_checks(frame)
    payload = build_audit_report(checked, build_summary(checked))
    worksheet = load_workbook(BytesIO(payload))["Flagged Entries"]

    description_cell = worksheet.cell(row=2, column=5)
    assert description_cell.value.startswith("=")
    assert description_cell.data_type == "s"


def test_compute_accuracy_hand_calculated(make_ledger) -> None:
    frame = make_ledger([{} for _ in range(5)])
    frame[GROUND_TRUTH_COLUMN] = ["Duplicate", "Clean", "Weekend", "Clean", "Duplicate"]
    for column in FLAG_COLUMNS:
        frame[column] = False
    frame.loc[[0, 4], "flag_duplicate"] = True
    frame.loc[1, "flag_round"] = True
    frame["flag_count"] = frame[list(FLAG_COLUMNS)].sum(axis=1).astype(int)

    accuracy = compute_accuracy(frame)

    assert accuracy["overall"] == {
        "tp": 2,
        "fp": 1,
        "fn": 1,
        "tn": 1,
        "precision": pytest.approx(2 / 3),
        "recall": pytest.approx(2 / 3),
        "f1": pytest.approx(2 / 3),
        "accuracy": pytest.approx(3 / 5),
    }
    assert accuracy["injected"] == 3
    assert accuracy["predicted"] == 3
    assert len(accuracy["missed"]) == 1
    assert len(accuracy["false_alarms"]) == 1

    per_type = {item["type"]: item for item in accuracy["per_type"]}
    assert set(per_type) == set(GT_TO_FLAG) | {"MultiRule"}
    assert per_type["Duplicate"]["tp"] == 2
    assert per_type["Duplicate"]["precision"] == 1.0
    assert per_type["Duplicate"]["recall"] == 1.0
    assert per_type["Weekend"]["fn"] == 1
    assert per_type["Weekend"]["recall"] == 0.0
    assert per_type["RoundAmount"]["fp"] == 1
    assert per_type["RoundAmount"]["precision"] == 0.0
    assert per_type["MultiRule"]["support"] == 0


def test_multi_rule_rows_require_at_least_two_flags(make_ledger) -> None:
    frame = make_ledger([{} for _ in range(3)])
    frame[GROUND_TRUTH_COLUMN] = ["MultiRule", "MultiRule", "Clean"]
    for column in FLAG_COLUMNS:
        frame[column] = False
    frame.loc[0, "flag_self_approval"] = True
    frame.loc[0, "flag_threshold"] = True
    frame.loc[1, "flag_self_approval"] = True
    frame["flag_count"] = frame[list(FLAG_COLUMNS)].sum(axis=1).astype(int)

    accuracy = compute_accuracy(frame)

    assert accuracy["overall"]["tp"] == 1
    assert accuracy["overall"]["fn"] == 1
    assert accuracy["overall"]["fp"] == 0
    assert accuracy["per_type"][-1]["type"] == "MultiRule"
    assert accuracy["per_type"][-1]["tp"] == 1
    assert accuracy["per_type"][-1]["recall"] == 0.5


def test_compute_accuracy_matches_the_demo(checked_demo: pd.DataFrame) -> None:
    accuracy = compute_accuracy(checked_demo)

    assert accuracy["overall"]["tp"] == 57
    assert accuracy["overall"]["fp"] == 0
    assert accuracy["overall"]["fn"] == 0
    assert accuracy["overall"]["tn"] == 549
    assert accuracy["overall"]["precision"] == 1.0
    assert accuracy["overall"]["recall"] == 1.0
    assert accuracy["overall"]["f1"] == 1.0
    assert accuracy["overall"]["accuracy"] == 1.0
    assert accuracy["tested_entries"] == 606
    assert len(accuracy["missed"]) == 0
    assert len(accuracy["false_alarms"]) == 0

    per_type = {item["type"]: item for item in accuracy["per_type"]}
    assert set(per_type) == set(GT_TO_FLAG) | {"MultiRule"}
    expected_support = {
        "Duplicate": 12,
        "RoundAmount": 5,
        "Weekend": 5,
        "JustBelowLimit": 6,
        "SelfApproval": 5,
        "BenfordManip": 18,
        "MultiRule": 6,
    }
    for label, support in expected_support.items():
        assert per_type[label]["support"] == support
        assert per_type[label]["tp"] == support
        assert per_type[label]["fp"] == 0
        assert per_type[label]["fn"] == 0
        assert per_type[label]["precision"] == 1.0
        assert per_type[label]["recall"] == 1.0

    assert accuracy["injected"] == 57
    assert accuracy["predicted"] == 57


def test_compute_accuracy_requires_ground_truth(checked_demo: pd.DataFrame) -> None:
    assert compute_accuracy(checked_demo.drop(columns=[GROUND_TRUTH_COLUMN])) is None


def test_dataframe_export_styles_and_round_trips(make_ledger) -> None:
    frame = make_ledger(
        [
            {"Amount": 1_234.56, "Description": "=1+1"},
            {"Amount": 2_345.67, "Description": "Office supplies"},
        ]
    )
    payload = dataframe_to_excel_bytes(frame, sheet_name="Bad/Name:*?[]")
    workbook = load_workbook(BytesIO(payload))
    worksheet = workbook["BadName"]

    assert workbook.sheetnames == ["BadName"]
    assert worksheet["A1"].fill.start_color.rgb.endswith("1F1B3A")
    assert worksheet["A1"].font.bold is True
    assert worksheet["B2"].number_format == "mm/dd/yyyy"
    assert worksheet["F2"].number_format == '"Rs. "#,##0.00'
    assert worksheet.column_dimensions["E"].width >= 11
    assert worksheet.freeze_panes == "A2"
    assert worksheet.sheet_view.showGridLines is False

    description_cell = worksheet["E2"]
    assert description_cell.value == "=1+1"
    assert description_cell.data_type == "s"

    round_trip = pd.read_excel(BytesIO(payload), sheet_name="BadName")
    assert list(round_trip.columns) == list(frame.columns)
    assert len(round_trip) == len(frame)
